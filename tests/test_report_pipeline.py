"""
Testes de regressão para o pipeline de geração de relatório (HTML e DOCX).
Garante que a estrutura gerada mantenha paridade e integridade.
"""

import os
import unittest
from execution.parse_content import parse_document_content
from execution.generate_report import render_report_html
from execution.generate_docx import generate_report_docx

class TestReportPipeline(unittest.TestCase):
    def setUp(self):
        self.sample_md_path = "inputs/exemplo_vistoria.md"
        self.template_path = "templates/report_template.html"
        self.tmp_html = "/tmp/test_report_out.html"
        self.tmp_docx = "/tmp/test_report_out.docx"

    def tearDown(self):
        for p in (self.tmp_html, self.tmp_docx):
            if os.path.exists(p):
                os.remove(p)

    def test_parse_markdown_metadata(self):
        with open(self.sample_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        data = parse_document_content(content)
        
        self.assertEqual(data["titulo"], "Relatório Técnico de Fiscalização de Obras Públicas")
        self.assertIn("Secretaria", data.get("orgao_val", ""))
        self.assertIn("Carlos", data.get("autor_val", ""))
        self.assertTrue(len(data.get("secoes", [])) >= 4)

    def test_generate_html_output(self):
        with open(self.sample_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        data = parse_document_content(content)
        out = render_report_html(data, self.template_path, self.tmp_html)
        self.assertTrue(os.path.exists(out))
        with open(out, "r", encoding="utf-8") as f:
            html = f.read()
        self.assertIn("Relatório Técnico de Fiscalização", html)
        self.assertIn("OBJETIVO DA VISTORIA", html)
        self.assertIn("<table", html)

    def test_generate_docx_output(self):
        with open(self.sample_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        data = parse_document_content(content)
        out = generate_report_docx(data, self.tmp_docx)
        self.assertTrue(os.path.exists(out))
        self.assertTrue(os.path.getsize(out) > 1000)

    def test_parse_docx_roundtrip(self):
        from execution.parse_docx import parse_docx_file
        with open(self.sample_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        data = parse_document_content(content)
        generate_report_docx(data, self.tmp_docx)
        
        # Lê o docx de volta
        model = parse_docx_file(self.tmp_docx)
        self.assertEqual(model.titulo, "RELATÓRIO TÉCNICO DE FISCALIZAÇÃO DE OBRAS PÚBLICAS")
        self.assertIn("Secretaria", model.orgao_val)
        self.assertIn("Carlos", model.autor_val)
        self.assertEqual(len(model.secoes), 4)

    def test_docx_image_extraction(self):
        import docx, io, base64
        from execution.parse_docx import parse_docx_file
        doc = docx.Document()
        png_1px = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==')
        p = doc.add_paragraph('1. SEÇÃO COM FOTO')
        r = p.add_run()
        r.add_picture(io.BytesIO(png_1px))
        tmp_img_docx = "/tmp/test_unit_img.docx"
        doc.save(tmp_img_docx)

        model = parse_docx_file(tmp_img_docx)
        self.assertTrue(any(b.type == "image" for s in model.secoes for b in s.blocks))
        if os.path.exists(tmp_img_docx):
            os.remove(tmp_img_docx)

    def test_parse_pdf_file(self):
        from execution.parse_pdf import parse_pdf_file
        if os.path.exists("/tmp/teste_doc.pdf"):
            model = parse_pdf_file("/tmp/teste_doc.pdf")
            self.assertTrue(len(model.titulo) > 0)
            self.assertTrue(len(model.secoes) >= 1)

if __name__ == "__main__":
    unittest.main()
